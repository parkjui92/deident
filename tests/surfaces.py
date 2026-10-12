"""숨은 자리 합성 코퍼스.

화면에 보이지 않는 자리(셀 메모·정의된 이름·발표자 노트·차트에 딸린 엑셀·PDF 주석 …)
에 합성 값을 심은 문서를 **자리 하나당 파일 하나**로 만든다. 치환기가 그 자리를
모르면 값이 산출물에 남거나 게이트가 막는다. 어느 쪽이든 이 코퍼스가 드러낸다.

판정은 도구 자신의 게이트에 맡기지 않는다. `residue()` 가 산출물을 따로 펼쳐
(중첩 zip·PDF 주석·양식·책갈피까지) 값이 남았는지 직접 찾는다. 게이트와 같은
코드로 채점하면 게이트가 못 보는 자리는 영원히 '통과'로 나온다.

값은 전부 허구다 (synth.py 와 같은 원칙).

    python3 tests/surfaces.py          # 형식 × 자리 표를 찍는다
"""

from __future__ import annotations

import html
import io
import re
import shutil
import sys
import tempfile
import unicodedata
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from synth import PLANTED  # noqa: E402

# 표에도 나오는 값 — 대장(매핑)에 오른다. 숨은 자리에 있어도 반드시 지워져야 한다.
KNOWN = PLANTED["person_pi"]
# 숨은 자리에만 있는 값. 번호·이메일은 모양만으로 잡혀야 한다.
HIDDEN_PHONE = "010-8642-9753"
HIDDEN_EMAIL = "seulgi.hb@krict.re.kr"
# 숨은 자리에만 있는 이름 — 직함 문맥이 붙어 있다. 놓쳐도 '탐지 한계'지만 따로 센다.
HIDDEN_NAME = "황보슬기"

NOTE = f"담당 {HIDDEN_NAME} 과장 {HIDDEN_PHONE} (책임 {KNOWN})"

KOREAN_FONTS = [
    "/System/Library/Fonts/Supplemental/AppleGothic.ttf",
    "/System/Library/Fonts/AppleSDGothicNeo.ttc",
    "/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
]

def _tiny_png() -> bytes:
    """1×1 PNG — 그림 자리를 만들기 위한 최소 파일."""
    import struct
    import zlib

    def chunk(kind: bytes, data: bytes) -> bytes:
        body = kind + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(b"\x00\xff\xff\xff")) + chunk(b"IEND", b""))


@dataclass
class Case:
    fmt: str
    surface: str
    build: Callable[[Path], None]
    must_remove: list[str]                      # 산출물에 남으면 안 되는 값
    soft: list[str] = field(default_factory=list)   # 남으면 '탐지 한계'로 따로 세는 값
    label: str = ""                             # 사람이 읽는 자리 이름

    @property
    def id(self) -> str:
        return f"{self.fmt}:{self.surface}"

    @property
    def filename(self) -> str:
        return f"{self.surface}.{self.fmt}"


# ── zip 손질 ────────────────────────────────────────────────────────────────
def patch_zip(path: Path, edits: dict[str, Callable[[bytes], bytes]] | None = None,
              add: dict[str, bytes] | None = None) -> None:
    """멤버를 고치거나 덧붙인다. 라이브러리가 만들어 주지 않는 자리를 심을 때 쓴다."""
    edits, add = edits or {}, add or {}
    buf = io.BytesIO()
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info)
            if info.filename in edits:
                data = edits[info.filename](data)
            if info.filename in add:
                continue
            dst.writestr(info, data)
        for name, data in add.items():
            dst.writestr(name, data)
    path.write_bytes(buf.getvalue())


def _x(text: str) -> str:
    return html.escape(text, quote=True)


# ── 엑셀 ────────────────────────────────────────────────────────────────────
def _xlsx_base():
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "참여연구원"
    ws.append(["연번", "성명", "소속", "연락처"])
    ws.append([1, KNOWN, PLANTED["org_a"], PLANTED["phone_pi"]])
    ws.append([2, PLANTED["person_co"], PLANTED["org_b"], "010-9876-5432"])
    return wb, ws


RRN_DIGITS = PLANTED["rrn_pi"].replace("-", "")
PHONE_AS_NUMBER = "1086429753"          # 숫자 셀에 넣으면 앞자리 0이 떨어진다


def _xlsx_numeric_rrn(path: Path) -> None:
    wb, ws = _xlsx_base()
    ws["E1"] = "주민등록번호"
    ws["E2"] = int(RRN_DIGITS)
    wb.save(path)


def _xlsx_numeric_phone(path: Path) -> None:
    wb, ws = _xlsx_base()
    ws["D3"] = int(PHONE_AS_NUMBER)
    wb.save(path)


def _xlsx_cell_comment(path: Path) -> None:
    from openpyxl.comments import Comment

    wb, ws = _xlsx_base()
    ws["B2"].comment = Comment(NOTE, KNOWN)
    wb.save(path)


def _xlsx_sheet_name(path: Path) -> None:
    wb, _ = _xlsx_base()
    sheet = wb.create_sheet(f"{KNOWN} 면담")
    sheet.append(["비고", "없음"])
    wb.save(path)


def _xlsx_defined_name(path: Path) -> None:
    from openpyxl.workbook.defined_name import DefinedName

    wb, _ = _xlsx_base()
    name = f"{KNOWN}_연락처"
    wb.defined_names[name] = DefinedName(name, attr_text="참여연구원!$D$2")
    wb.save(path)


def _xlsx_header_footer(path: Path) -> None:
    wb, ws = _xlsx_base()
    ws.oddHeader.center.text = NOTE
    ws.oddFooter.left.text = f"작성 {KNOWN}"
    wb.save(path)


def _xlsx_hyperlink(path: Path) -> None:
    wb, ws = _xlsx_base()
    ws["E2"] = "메일 보내기"
    ws["E2"].hyperlink = f"mailto:{HIDDEN_EMAIL}"
    ws["E2"].hyperlink.tooltip = NOTE
    wb.save(path)


def _xlsx_formula_cached(path: Path) -> None:
    wb, ws = _xlsx_base()
    ws["E1"] = "표시"
    ws["E2"] = '=B2&" "&D2'
    wb.save(path)
    cached = f"{KNOWN} {PLANTED['phone_pi']}"

    def edit(data: bytes) -> bytes:
        text = data.decode("utf-8")
        text, n = re.subn(r'<c r="E2"([^>]*)><f>(.*?)</f><v\s*/>',
                          lambda m: f'<c r="E2"{m.group(1)} t="str"><f>{m.group(2)}</f>'
                                    f"<v>{_x(cached)}</v>", text)
        if n == 0:
            text, n = re.subn(r'<c r="E2"([^>]*)><f>(.*?)</f><v></v>',
                              lambda m: f'<c r="E2"{m.group(1)} t="str"><f>{m.group(2)}</f>'
                                        f"<v>{_x(cached)}</v>", text)
        assert n == 1, "수식 셀을 찾지 못했다"
        return text.encode("utf-8")

    patch_zip(path, edits={"xl/worksheets/sheet1.xml": edit})


def _xlsx_formula_literal(path: Path) -> None:
    wb, ws = _xlsx_base()
    ws["E1"] = "구분"
    ws["E2"] = f'=IF(B2="{KNOWN}","책임","공동")'
    wb.save(path)


def _xlsx_data_validation(path: Path) -> None:
    from openpyxl.worksheet.datavalidation import DataValidation

    wb, ws = _xlsx_base()
    dv = DataValidation(type="list", formula1=f'"{KNOWN},{PLANTED["person_co"]}"',
                        allow_blank=True)
    dv.prompt = NOTE
    dv.promptTitle = "담당자 선택"
    ws.add_data_validation(dv)
    dv.add("F2:F10")
    wb.save(path)


def _xlsx_docprops(path: Path) -> None:
    wb, _ = _xlsx_base()
    wb.properties.creator = KNOWN
    wb.properties.lastModifiedBy = KNOWN
    wb.properties.title = NOTE
    wb.properties.description = NOTE
    wb.properties.keywords = HIDDEN_NAME
    wb.save(path)


def _xlsx_drawing_textbox(path: Path) -> None:
    wb, _ = _xlsx_base()
    wb.save(path)
    drawing = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        "<xdr:twoCellAnchor><xdr:sp><xdr:nvSpPr>"
        f'<xdr:cNvPr id="2" name="TextBox 1" descr="{_x(NOTE)}"/><xdr:cNvSpPr txBox="1"/>'
        "</xdr:nvSpPr><xdr:txBody><a:bodyPr/><a:p><a:r><a:t>"
        f"{_x(NOTE)}</a:t></a:r></a:p></xdr:txBody></xdr:sp></xdr:twoCellAnchor></xdr:wsDr>"
    ).encode("utf-8")
    patch_zip(path, add={"xl/drawings/drawing1.xml": drawing})


def _xlsx_threaded_comment(path: Path) -> None:
    wb, _ = _xlsx_base()
    wb.save(path)
    ns = "http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments"
    persons = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<personList xmlns="{ns}"><person displayName="{_x(KNOWN)}" id="{{A1}}" '
        f'userId="{_x(HIDDEN_EMAIL)}" providerId="AD"/></personList>').encode("utf-8")
    thread = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<ThreadedComments xmlns="{ns}"><threadedComment ref="B2" personId="{{A1}}" id="{{C1}}">'
        f"<text>{_x(NOTE)}</text></threadedComment></ThreadedComments>").encode("utf-8")
    patch_zip(path, add={"xl/persons/person.xml": persons,
                         "xl/threadedComments/threadedComment1.xml": thread})


def _xlsx_pivot_cache(path: Path) -> None:
    wb, _ = _xlsx_base()
    wb.save(path)
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    cache = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<pivotCacheDefinition xmlns="{ns}" refreshedBy="{_x(KNOWN)}"><cacheFields count="2">'
        f'<cacheField name="성명"><sharedItems><s v="{_x(KNOWN)}"/><s v="{_x(HIDDEN_NAME)}"/>'
        f'</sharedItems></cacheField><cacheField name="연락처"><sharedItems>'
        f'<s v="{HIDDEN_PHONE}"/></sharedItems></cacheField></cacheFields>'
        "</pivotCacheDefinition>").encode("utf-8")
    patch_zip(path, add={"xl/pivotCache/pivotCacheDefinition1.xml": cache})


def _xlsx_table_def(path: Path) -> None:
    wb, _ = _xlsx_base()
    wb.save(path)
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    table = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<table xmlns="{ns}" id="1" name="{_x(KNOWN)}_표" displayName="{_x(KNOWN)}_표" '
        f'ref="A1:D3"><tableColumns count="1"><tableColumn id="1" name="{_x(NOTE)}"/>'
        "</tableColumns></table>").encode("utf-8")
    patch_zip(path, add={"xl/tables/table1.xml": table})


# ── 워드 ────────────────────────────────────────────────────────────────────
def _docx_base():
    import docx

    doc = docx.Document()
    doc.add_paragraph("연구개발계획서 (합성 시험본)")
    table = doc.add_table(rows=2, cols=3)
    for col, text in enumerate(["성명", "소속", "연락처"]):
        table.cell(0, col).text = text
    for col, text in enumerate([KNOWN, PLANTED["org_a"], PLANTED["phone_pi"]]):
        table.cell(1, col).text = text
    doc.add_paragraph("위 연구책임자가 과제를 총괄한다.")
    return doc


def _docx_header_footer(path: Path) -> None:
    doc = _docx_base()
    section = doc.sections[0]
    section.header.paragraphs[0].text = NOTE
    section.footer.paragraphs[0].text = f"작성 {KNOWN}"
    doc.save(path)


def _docx_comment(path: Path) -> None:
    doc = _docx_base()
    para = doc.paragraphs[-1]
    doc.add_comment(runs=para.runs, text=NOTE, author=KNOWN, initials="K")
    doc.save(path)


def _docx_core_props(path: Path) -> None:
    doc = _docx_base()
    props = doc.core_properties
    props.author = KNOWN
    props.last_modified_by = KNOWN
    props.title = NOTE
    props.subject = NOTE
    props.keywords = HIDDEN_NAME
    props.category = HIDDEN_NAME
    props.comments = NOTE
    doc.save(path)


def _docx_hyperlink(path: Path) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = _docx_base()
    rid = doc.part.relate_to(f"mailto:{HIDDEN_EMAIL}", RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    link.set(qn("w:tooltip"), NOTE)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "문의하기"
    run.append(text)
    link.append(run)
    doc.add_paragraph()._p.append(link)
    doc.save(path)


def _docx_picture_alt(path: Path) -> None:
    doc = _docx_base()
    doc.add_picture(io.BytesIO(_tiny_png()))
    for node in doc.element.iter():
        if node.tag.endswith("}docPr"):
            node.set("name", f"{KNOWN} 서명")
            node.set("descr", NOTE)
            node.set("title", NOTE)
    doc.save(path)


def _docx_bookmark(path: Path) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = _docx_base()
    para = doc.paragraphs[-1]._p
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), "1")
    start.set(qn("w:name"), f"{KNOWN}_서명")
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), "1")
    para.insert(0, start)
    para.append(end)
    doc.save(path)


def _docx_doc_var(path: Path) -> None:
    doc = _docx_base()
    doc.save(path)

    def edit(data: bytes) -> bytes:
        text = data.decode("utf-8")
        var = f'<w:docVars><w:docVar w:name="담당자" w:val="{_x(NOTE)}"/></w:docVars>'
        assert "</w:settings>" in text
        return text.replace("</w:settings>", var + "</w:settings>").encode("utf-8")

    patch_zip(path, edits={"word/settings.xml": edit})


def _docx_content_control(path: Path) -> None:
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    doc = _docx_base()
    sdt = parse_xml(
        f"<w:sdt {nsdecls('w')}><w:sdtPr>"
        f'<w:alias w:val="{_x(NOTE)}"/><w:tag w:val="{_x(KNOWN)}"/>'
        "</w:sdtPr><w:sdtContent><w:p><w:r><w:t>"
        f"{_x(NOTE)}</w:t></w:r></w:p></w:sdtContent></w:sdt>")
    doc.element.body.insert(len(doc.element.body) - 1, sdt)
    doc.save(path)


def _docx_custom_xml(path: Path) -> None:
    doc = _docx_base()
    doc.save(path)
    item = (f'<?xml version="1.0" encoding="UTF-8"?><과제><담당>{_x(NOTE)}</담당></과제>'
            ).encode("utf-8")
    patch_zip(path, add={"customXml/item1.xml": item})


def _embedded_workbook() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["성명", "건수", "비고"])
    ws.append([KNOWN, 3, ""])
    ws.append([HIDDEN_NAME, 5, NOTE])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _chart_xml() -> bytes:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><c:chart>'
        f"<c:title><c:tx><c:rich><a:p><a:r><a:t>{_x(NOTE)}</a:t></a:r></a:p></c:rich></c:tx></c:title>"
        "<c:plotArea><c:barChart><c:ser><c:cat><c:strRef><c:f>Sheet1!$A$2:$A$3</c:f><c:strCache>"
        f'<c:pt idx="0"><c:v>{_x(KNOWN)}</c:v></c:pt><c:pt idx="1"><c:v>{_x(HIDDEN_NAME)}</c:v></c:pt>'
        "</c:strCache></c:strRef></c:cat></c:ser></c:barChart></c:plotArea></c:chart></c:chartSpace>"
    ).encode("utf-8")


def _docx_chart(path: Path) -> None:
    doc = _docx_base()
    doc.save(path)
    patch_zip(path, add={
        "word/charts/chart1.xml": _chart_xml(),
        "word/embeddings/Microsoft_Excel_Sheet1.xlsx": _embedded_workbook(),
    })


def _docx_field(path: Path) -> None:
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    doc = _docx_base()
    field_xml = parse_xml(
        f'<w:p {nsdecls("w")}><w:fldSimple w:instr=" HYPERLINK &quot;mailto:{HIDDEN_EMAIL}&quot; ">'
        f"<w:r><w:t>{_x(KNOWN)}</w:t></w:r></w:fldSimple></w:p>")
    doc.element.body.insert(len(doc.element.body) - 1, field_xml)
    doc.save(path)


# ── 발표자료 ────────────────────────────────────────────────────────────────
def _pptx_base():
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "과제 추진체계 (합성 시험본)"
    frame = slide.shapes.add_table(2, 3, Inches(0.5), Inches(1.5), Inches(8), Inches(1))
    table = frame.table
    for col, text in enumerate(["성명", "소속", "연락처"]):
        table.cell(0, col).text = text
    for col, text in enumerate([KNOWN, PLANTED["org_a"], PLANTED["phone_pi"]]):
        table.cell(1, col).text = text
    return prs, slide


def _pptx_notes(path: Path) -> None:
    prs, slide = _pptx_base()
    slide.notes_slide.notes_text_frame.text = NOTE
    prs.save(path)


def _pptx_shape_alt(path: Path) -> None:
    from pptx.util import Inches

    prs, slide = _pptx_base()
    box = slide.shapes.add_textbox(Inches(0.5), Inches(4), Inches(3), Inches(0.5))
    box.text_frame.text = "서명란"
    box.name = f"{KNOWN} 서명"
    box._element.nvSpPr.cNvPr.set("descr", NOTE)
    prs.save(path)


def _pptx_comment(path: Path) -> None:
    prs, _ = _pptx_base()
    prs.save(path)
    ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    authors = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<p:cmAuthorLst xmlns:p="{ns}"><p:cmAuthor id="0" name="{_x(KNOWN)}" initials="K" '
        'lastIdx="1" clrIdx="0"/></p:cmAuthorLst>').encode("utf-8")
    comments = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<p:cmLst xmlns:p="{ns}"><p:cm authorId="0" dt="2026-01-01T00:00:00.000" idx="1">'
        f'<p:pos x="10" y="10"/><p:text>{_x(NOTE)}</p:text></p:cm></p:cmLst>').encode("utf-8")
    patch_zip(path, add={"ppt/commentAuthors.xml": authors,
                         "ppt/comments/comment1.xml": comments})


def _pptx_chart(path: Path) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs, slide = _pptx_base()
    data = CategoryChartData()
    data.categories = [KNOWN, HIDDEN_NAME]
    data.add_series(f"{KNOWN} 팀 실적", (3, 5))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(3),
                           Inches(6), Inches(3), data)
    prs.save(path)


def _pptx_smartart(path: Path) -> None:
    prs, _ = _pptx_base()
    prs.save(path)
    a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    dgm = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
    dsp = "http://schemas.microsoft.com/office/drawing/2008/diagram"
    items = [f"{HIDDEN_NAME} 본부장", f"{KNOWN} 책임연구원", NOTE]
    pts = "".join(
        f'<dgm:pt modelId="{i + 1}"><dgm:prSet/><dgm:spPr/><dgm:t><a:bodyPr/><a:p><a:r>'
        f"<a:t>{_x(t)}</a:t></a:r></a:p></dgm:t></dgm:pt>" for i, t in enumerate(items))
    data = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<dgm:dataModel xmlns:dgm="{dgm}" xmlns:a="{a}"><dgm:ptLst>{pts}</dgm:ptLst>'
            "</dgm:dataModel>").encode("utf-8")
    shapes = "".join(
        f'<dsp:sp modelId="{i + 1}"><dsp:txBody><a:bodyPr/><a:p><a:r><a:t>{_x(t)}</a:t></a:r>'
        "</a:p></dsp:txBody></dsp:sp>" for i, t in enumerate(items))
    drawing = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
               f'<dsp:drawing xmlns:dsp="{dsp}" xmlns:a="{a}"><dsp:spTree>{shapes}</dsp:spTree>'
               "</dsp:drawing>").encode("utf-8")
    patch_zip(path, add={"ppt/diagrams/data1.xml": data,
                         "ppt/diagrams/drawing1.xml": drawing})


def _pptx_core_props(path: Path) -> None:
    prs, _ = _pptx_base()
    props = prs.core_properties
    props.author = KNOWN
    props.last_modified_by = KNOWN
    props.title = NOTE
    props.subject = NOTE
    props.keywords = HIDDEN_NAME
    props.comments = NOTE
    prs.save(path)


def _pptx_hyperlink(path: Path) -> None:
    from pptx.util import Inches

    prs, slide = _pptx_base()
    box = slide.shapes.add_textbox(Inches(0.5), Inches(4), Inches(3), Inches(0.5))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "문의하기"
    run.hyperlink.address = f"mailto:{HIDDEN_EMAIL}"
    prs.save(path)


def _pptx_section_name(path: Path) -> None:
    prs, _ = _pptx_base()
    prs.save(path)

    def edit(data: bytes) -> bytes:
        text = data.decode("utf-8")
        ext = ('<p:extLst><p:ext uri="{521415D9-36F7-43E2-AB2F-B90AF26B5E84}">'
               '<p14:sectionLst xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main">'
               f'<p14:section name="{_x(KNOWN)} 발표" id="{{00000000-0000-0000-0000-000000000001}}">'
               "<p14:sldIdLst/></p14:section></p14:sectionLst></p:ext></p:extLst>")
        assert "</p:presentation>" in text
        return text.replace("</p:presentation>", ext + "</p:presentation>").encode("utf-8")

    patch_zip(path, edits={"ppt/presentation.xml": edit})


# ── PDF ─────────────────────────────────────────────────────────────────────
def korean_font() -> str | None:
    for candidate in KOREAN_FONTS:
        if Path(candidate).exists():
            return candidate
    return None


def _pdf_base():
    import fitz

    font = korean_font()
    if font is None:
        raise RuntimeError("한글 글꼴이 없다")
    doc = fitz.open()
    page = doc.new_page()
    page.insert_font(fontname="KR", fontfile=font)
    y = 80
    for line in ["연구개발계획서 (합성 시험본)", f"연구책임자: {KNOWN}",
                 f"연락처: {PLANTED['phone_pi']}"]:
        page.insert_text((60, y), line, fontname="KR", fontsize=12)
        y += 26
    return doc, page


def _pdf_annot_note(path: Path) -> None:
    doc, page = _pdf_base()
    annot = page.add_text_annot((300, 100), NOTE)
    annot.set_info(title=KNOWN, subject=NOTE)
    annot.update()
    doc.save(path)
    doc.close()


def _pdf_annot_freetext(path: Path) -> None:
    import fitz

    doc, page = _pdf_base()
    page.add_freetext_annot(fitz.Rect(60, 300, 400, 340),
                            f"contact {HIDDEN_PHONE} {HIDDEN_EMAIL}")
    doc.save(path)
    doc.close()


def _pdf_form_field(path: Path) -> None:
    import fitz

    doc, page = _pdf_base()
    widget = fitz.Widget()
    widget.field_type = fitz.PDF_WIDGET_TYPE_TEXT
    widget.field_name = "담당자"
    widget.field_value = NOTE
    widget.rect = fitz.Rect(60, 200, 400, 225)
    page.add_widget(widget)
    doc.save(path)
    doc.close()


def _pdf_bookmark(path: Path) -> None:
    doc, _ = _pdf_base()
    doc.set_toc([[1, f"{KNOWN} 이력 {HIDDEN_PHONE}", 1]])
    doc.save(path)
    doc.close()


def _pdf_link(path: Path) -> None:
    import fitz

    doc, page = _pdf_base()
    page.insert_link({"kind": fitz.LINK_URI, "from": fitz.Rect(60, 250, 200, 270),
                      "uri": f"mailto:{HIDDEN_EMAIL}"})
    doc.save(path)
    doc.close()


def _pdf_attachment(path: Path) -> None:
    doc, _ = _pdf_base()
    doc.embfile_add("명단.txt", NOTE.encode("utf-8"), desc=NOTE)
    doc.save(path)
    doc.close()


def _pdf_metadata(path: Path) -> None:
    doc, _ = _pdf_base()
    doc.set_metadata({"author": KNOWN, "title": NOTE, "subject": NOTE,
                      "keywords": HIDDEN_NAME})
    doc.save(path)
    doc.close()


def _pdf_hidden_text(path: Path) -> None:
    doc, page = _pdf_base()
    page.insert_text((60, 400), NOTE, fontname="KR", fontsize=12, render_mode=3)
    doc.save(path)
    doc.close()



# ── 한글 ────────────────────────────────────────────────────────────────────
HP_NS = "http://www.hancom.co.kr/hwpml/2011/paragraph"
HS_NS = "http://www.hancom.co.kr/hwpml/2011/section"


def _hp_para(text: str, inner: str = "") -> str:
    return f"<hp:p><hp:run>{inner}<hp:t>{_x(text)}</hp:t></hp:run></hp:p>"


def _hwpx_build(path: Path, body: str = "", meta: str = "", preview: str = "") -> None:
    """추출기가 읽는 뼈대만 갖춘 최소 hwpx (한글에서 열리는 완전한 파일은 아니다)."""
    def cell(row: int, col: int, text: str) -> str:
        return (f"<hp:tc><hp:subList>{_hp_para(text)}</hp:subList>"
                f'<hp:cellAddr colAddr="{col}" rowAddr="{row}"/>'
                '<hp:cellSpan colSpan="1" rowSpan="1"/></hp:tc>')

    rows = [["성명", "소속", "연락처"], [KNOWN, PLANTED["org_a"], PLANTED["phone_pi"]]]
    table = '<hp:tbl rowCnt="2" colCnt="3">' + "".join(
        "<hp:tr>" + "".join(cell(r, c, text) for c, text in enumerate(row)) + "</hp:tr>"
        for r, row in enumerate(rows)) + "</hp:tbl>"
    section = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<hs:sec xmlns:hs="{HS_NS}" xmlns:hp="{HP_NS}">'
        + _hp_para("연구개발계획서 (합성 시험본)")
        + f"<hp:p><hp:run>{table}</hp:run></hp:p>"
        + _hp_para("위 연구책임자가 과제를 총괄한다.") + body + "</hs:sec>")
    content = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<opf:package xmlns:opf="http://www.idpf.org/2007/opf/">'
        f'<opf:metadata><opf:meta name="generator" content="synthetic"/>{meta}</opf:metadata>'
        '<opf:manifest><opf:item id="section0" href="Contents/section0.xml" '
        'media-type="application/xml"/></opf:manifest></opf:package>')
    header = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
              '<hh:head xmlns:hh="http://www.hancom.co.kr/hwpml/2011/head" version="1.4">'
              '<hh:refList><hh:fontfaces><hh:fontface lang="HANGUL"><hh:font id="0" '
              'face="함초롬바탕" type="TTF"/></hh:fontface></hh:fontfaces>'
              '<hh:styles><hh:style id="0" type="PARA" name="바탕글" engName="Normal"/>'
              "</hh:styles></hh:refList></hh:head>")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("mimetype", "application/hwp+zip")
        zf.writestr("Contents/content.hpf", content)
        zf.writestr("Contents/header.xml", header)
        zf.writestr("Contents/section0.xml", section)
        zf.writestr("Preview/PrvText.txt", preview or "연구개발계획서")


def _hwpx_bookmark(path: Path) -> None:
    _hwpx_build(path, _hp_para(
        "서명란", f'<hp:ctrl><hp:bookmark name="{_x(KNOWN)}_서명"/></hp:ctrl>'))


def _hwpx_hyperlink(path: Path) -> None:
    field = ('<hp:ctrl><hp:fieldBegin id="1" type="HYPERLINK" name="" editable="0">'
             '<hp:parameters cnt="1" name=""><hp:stringParam name="Command">'
             f"mailto:{HIDDEN_EMAIL};1;0;0;</hp:stringParam></hp:parameters>"
             "</hp:fieldBegin></hp:ctrl>")
    _hwpx_build(path, _hp_para(f"{KNOWN} 책임에게 문의", field))


def _hwpx_click_here(path: Path) -> None:
    field = (f'<hp:ctrl><hp:fieldBegin id="2" type="CLICK_HERE" name="{_x(NOTE)}" editable="1">'
             '<hp:parameters cnt="1" name=""><hp:stringParam name="Direction">'
             f"{_x(NOTE)}</hp:stringParam></hp:parameters></hp:fieldBegin></hp:ctrl>")
    _hwpx_build(path, _hp_para("담당자를 적으십시오", field))


def _hwpx_form_list(path: Path) -> None:
    combo = ('<hp:comboBox name="담당자 선택" listBoxRows="4">'
             f'<hp:listItem displayText="{_x(KNOWN)}" value="{HIDDEN_PHONE}"/>'
             f'<hp:listItem displayText="{_x(HIDDEN_NAME)} 과장" value="{HIDDEN_EMAIL}"/>'
             "</hp:comboBox>")
    _hwpx_build(path, _hp_para("담당자", combo))


def _hwpx_shape_comment(path: Path) -> None:
    pic = f"<hp:pic><hp:shapeComment>{_x(NOTE)}</hp:shapeComment></hp:pic>"
    _hwpx_build(path, _hp_para("서명 그림", pic))


def _hwpx_metadata(path: Path) -> None:
    _hwpx_build(path, meta=(
        f'<opf:meta name="creator" content="{_x(KNOWN)}"/>'
        f'<opf:meta name="title" content="{_x(NOTE)}"/>'
        f"<opf:title>{_x(NOTE)}</opf:title>"), preview=NOTE)


NUMS = [HIDDEN_PHONE]
ALL = [KNOWN, HIDDEN_PHONE]

CASES: list[Case] = [
    # 엑셀
    Case("xlsx", "numeric_rrn", _xlsx_numeric_rrn, [KNOWN, RRN_DIGITS], label="숫자 셀(주민등록번호)"),
    Case("xlsx", "numeric_phone", _xlsx_numeric_phone, [KNOWN, PHONE_AS_NUMBER], label="숫자 셀(연락처)"),
    Case("xlsx", "cell_comment", _xlsx_cell_comment, ALL, [HIDDEN_NAME], "셀 메모"),
    Case("xlsx", "threaded_comment", _xlsx_threaded_comment, [KNOWN, HIDDEN_PHONE, HIDDEN_EMAIL],
         [HIDDEN_NAME], "스레드 댓글·작성자"),
    Case("xlsx", "sheet_name", _xlsx_sheet_name, [KNOWN], label="시트 이름"),
    Case("xlsx", "defined_name", _xlsx_defined_name, [KNOWN], label="정의된 이름"),
    Case("xlsx", "header_footer", _xlsx_header_footer, ALL, [HIDDEN_NAME], "인쇄 머리글·바닥글"),
    Case("xlsx", "hyperlink", _xlsx_hyperlink, [KNOWN, HIDDEN_PHONE, HIDDEN_EMAIL],
         [HIDDEN_NAME], "하이퍼링크 주소·도구 설명"),
    Case("xlsx", "formula_cached", _xlsx_formula_cached, [KNOWN, PLANTED["phone_pi"]],
         label="수식 결과(저장된 값)"),
    Case("xlsx", "formula_literal", _xlsx_formula_literal, [KNOWN], label="수식 안 글자"),
    Case("xlsx", "data_validation", _xlsx_data_validation, ALL, [HIDDEN_NAME], "목록 상자·입력 안내"),
    Case("xlsx", "docprops", _xlsx_docprops, ALL, [HIDDEN_NAME], "문서 속성"),
    Case("xlsx", "drawing_textbox", _xlsx_drawing_textbox, ALL, [HIDDEN_NAME], "도형 글상자"),
    Case("xlsx", "pivot_cache", _xlsx_pivot_cache, ALL, [HIDDEN_NAME], "피벗 캐시"),
    Case("xlsx", "table_def", _xlsx_table_def, ALL, [HIDDEN_NAME], "표 정의(이름·열 이름)"),
    # 워드
    Case("docx", "header_footer", _docx_header_footer, ALL, [HIDDEN_NAME], "머리말·꼬리말"),
    Case("docx", "comment", _docx_comment, ALL, [HIDDEN_NAME], "메모"),
    Case("docx", "core_props", _docx_core_props, ALL, [HIDDEN_NAME], "문서 속성"),
    Case("docx", "hyperlink", _docx_hyperlink, [KNOWN, HIDDEN_PHONE, HIDDEN_EMAIL],
         [HIDDEN_NAME], "하이퍼링크 주소·도구 설명"),
    Case("docx", "picture_alt", _docx_picture_alt, ALL, [HIDDEN_NAME], "그림 이름·대체 텍스트"),
    Case("docx", "bookmark", _docx_bookmark, [KNOWN], label="책갈피 이름"),
    Case("docx", "doc_var", _docx_doc_var, ALL, [HIDDEN_NAME], "문서 변수"),
    Case("docx", "content_control", _docx_content_control, ALL, [HIDDEN_NAME], "콘텐츠 컨트롤"),
    Case("docx", "custom_xml", _docx_custom_xml, ALL, [HIDDEN_NAME], "customXml"),
    Case("docx", "chart", _docx_chart, ALL, [HIDDEN_NAME], "차트·내장 엑셀"),
    Case("docx", "field", _docx_field, [KNOWN, HIDDEN_EMAIL], label="필드 지시문"),
    # 발표자료
    Case("pptx", "notes", _pptx_notes, ALL, [HIDDEN_NAME], "발표자 노트"),
    Case("pptx", "shape_alt", _pptx_shape_alt, ALL, [HIDDEN_NAME], "도형 이름·대체 텍스트"),
    Case("pptx", "comment", _pptx_comment, ALL, [HIDDEN_NAME], "메모·작성자"),
    Case("pptx", "chart", _pptx_chart, [KNOWN], [HIDDEN_NAME], "차트·내장 엑셀"),
    Case("pptx", "smartart", _pptx_smartart, ALL, [HIDDEN_NAME], "스마트아트(조직도)"),
    Case("pptx", "core_props", _pptx_core_props, ALL, [HIDDEN_NAME], "문서 속성"),
    Case("pptx", "hyperlink", _pptx_hyperlink, [KNOWN, HIDDEN_EMAIL], label="하이퍼링크 주소"),
    Case("pptx", "section_name", _pptx_section_name, [KNOWN], label="구역 이름"),
    # 한글
    Case("hwpx", "bookmark", _hwpx_bookmark, [KNOWN], label="책갈피 이름"),
    Case("hwpx", "hyperlink", _hwpx_hyperlink, [KNOWN, HIDDEN_EMAIL], label="하이퍼링크 주소"),
    Case("hwpx", "click_here", _hwpx_click_here, ALL, [HIDDEN_NAME], "누름틀 이름·안내문"),
    Case("hwpx", "form_list", _hwpx_form_list, [KNOWN, HIDDEN_PHONE, HIDDEN_EMAIL],
         [HIDDEN_NAME], "양식 목록 항목"),
    Case("hwpx", "shape_comment", _hwpx_shape_comment, ALL, [HIDDEN_NAME], "그림 설명"),
    Case("hwpx", "metadata", _hwpx_metadata, ALL, [HIDDEN_NAME], "문서 정보·미리보기"),
    # PDF
    Case("pdf", "annot_note", _pdf_annot_note, ALL, [HIDDEN_NAME], "주석(메모)·작성자"),
    Case("pdf", "annot_freetext", _pdf_annot_freetext, [KNOWN, HIDDEN_PHONE, HIDDEN_EMAIL],
         label="주석(글상자)"),
    Case("pdf", "form_field", _pdf_form_field, ALL, [HIDDEN_NAME], "양식 필드 값"),
    Case("pdf", "bookmark", _pdf_bookmark, ALL, label="책갈피"),
    Case("pdf", "link", _pdf_link, [KNOWN, HIDDEN_EMAIL], label="링크 주소"),
    Case("pdf", "attachment", _pdf_attachment, ALL, [HIDDEN_NAME], "첨부 파일"),
    Case("pdf", "metadata", _pdf_metadata, ALL, [HIDDEN_NAME], "문서 정보"),
    Case("pdf", "hidden_text", _pdf_hidden_text, ALL, [HIDDEN_NAME], "보이지 않는 글자"),
]


# ── 독립 판정 ───────────────────────────────────────────────────────────────
RE_SEP = re.compile(r"[\s\-.]")


def _pdf_texts(data: bytes) -> list[tuple[str, str]]:
    import fitz

    out: list[tuple[str, str]] = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        out.append(("metadata", " ".join(str(v) for v in (doc.metadata or {}).values() if v)))
        try:
            out.append(("xmp", doc.get_xml_metadata() or ""))
        except Exception:
            pass
        for level, title, *_ in doc.get_toc(simple=True):
            out.append(("toc", title))
        for name in doc.embfile_names():
            info = doc.embfile_info(name)
            out.append((f"embfile:{name}", " ".join(str(v) for v in info.values())))
            out.extend((f"embfile:{name}/{w}", t) for w, t in dump_texts_bytes(
                doc.embfile_get(name), name))
        for page in doc:
            out.append((f"page{page.number + 1}", page.get_text("text")))
            for annot in page.annots() or []:
                info = annot.info
                out.append((f"page{page.number + 1}/annot",
                            " ".join(str(v) for v in info.values())))
            for widget in page.widgets() or []:
                out.append((f"page{page.number + 1}/widget",
                            f"{widget.field_name} {widget.field_value} {widget.field_label}"))
            for link in page.get_links():
                out.append((f"page{page.number + 1}/link", str(link.get("uri", ""))))
        # 객체 사전과 스트림까지 — API 가 보여 주지 않는 곳에 남은 문자열을 찾는다
        for xref in range(1, doc.xref_length()):
            try:
                out.append((f"xref{xref}", doc.xref_object(xref, compressed=False)))
                if doc.xref_is_stream(xref):
                    stream = doc.xref_stream(xref) or b""
                    out.append((f"xref{xref}/stream", stream.decode("utf-8", "ignore")))
                    out.append((f"xref{xref}/stream16", stream.decode("utf-16-be", "ignore")))
            except Exception:
                continue
    return out


RE_PDF_HEX = re.compile(r"<(FEFF[0-9A-Fa-f]+)>")


def _unhex_pdf_strings(text: str) -> str:
    """PDF 사전 속 <FEFF…> 16진 문자열을 글자로 풀어 덧붙인다."""
    extra = []
    for m in RE_PDF_HEX.finditer(text):
        try:
            extra.append(bytes.fromhex(m.group(1)).decode("utf-16", "ignore"))
        except ValueError:
            continue
    return text + " " + " ".join(extra) if extra else text


def dump_texts_bytes(data: bytes, name: str) -> list[tuple[str, str]]:
    """바이트 덩어리를 읽을 수 있는 글자로 펼친다 (중첩 zip 은 다시 연다)."""
    out: list[tuple[str, str]] = []
    if data[:5] == b"%PDF-":
        try:
            return [(w, _unhex_pdf_strings(t)) for w, t in _pdf_texts(data)]
        except Exception:
            pass
    if data[:2] == b"PK" and zipfile.is_zipfile(io.BytesIO(data)):
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            for info in zf.infolist():
                out.append((f"{info.filename}#name", info.filename))
                if info.is_dir():
                    continue
                out.extend((f"{info.filename}>{w}" if w else info.filename, t)
                           for w, t in dump_texts_bytes(zf.read(info), info.filename))
        return out
    for encoding in ("utf-8", "utf-16-le"):
        text = data.decode(encoding, "ignore")
        if text.strip():
            out.append(("" if encoding == "utf-8" else encoding, html.unescape(text)))
    return out


def dump_texts(path: Path) -> list[tuple[str, str]]:
    return dump_texts_bytes(path.read_bytes(), path.name)


def residue(path: Path, value: str) -> list[str]:
    """값이 남아 있는 자리의 이름들. 비어 있으면 깨끗하다."""
    forms = {unicodedata.normalize("NFC", value), unicodedata.normalize("NFD", value)}
    digits = "".join(ch for ch in value if ch.isdigit())
    numeric = len(digits) >= 7 and len(digits) >= len(value) - 4
    where: list[str] = []
    for name, text in dump_texts(path):
        if any(form in text for form in forms):
            where.append(name or path.name)
        elif numeric and digits in RE_SEP.sub("", text):
            where.append(name or path.name)
    return sorted(set(where))


def opens(path: Path) -> str:
    """산출물을 그 형식의 라이브러리로 다시 열어 본다. 문제가 있으면 사유를 돌려준다."""
    suffix = path.suffix.lower()
    try:
        if suffix == ".xlsx":
            import openpyxl

            wb = openpyxl.load_workbook(path)
            for ws in wb:
                if any(ch in ws.title for ch in "[]:*?/\\") or len(ws.title) > 31:
                    return "엑셀이 받지 않는 시트 이름"
                for row in ws.iter_rows():
                    for cell in row:
                        _ = cell.value
        elif suffix == ".docx":
            import docx

            _ = [p.text for p in docx.Document(str(path)).paragraphs]
        elif suffix == ".pptx":
            from pptx import Presentation

            _ = [shape.name for slide in Presentation(str(path)).slides
                 for shape in slide.shapes]
        elif suffix == ".pdf":
            import fitz

            with fitz.open(path) as doc:
                _ = [page.get_text("text") for page in doc]
        elif suffix == ".hwpx":
            from lxml import etree

            with zipfile.ZipFile(path) as zf:
                for name in zf.namelist():
                    if name.endswith((".xml", ".hpf")):
                        etree.fromstring(zf.read(name))
    except Exception as exc:   # noqa: BLE001
        return f"{type(exc).__name__}"
    return ""


# ── 실행 ────────────────────────────────────────────────────────────────────
CLEAN, BLOCKED, LEAK, ERROR = "깨끗함", "게이트 차단", "조용히 남음", "오류"


@dataclass
class Outcome:
    case: Case
    status: str
    detail: list[str] = field(default_factory=list)
    soft_left: list[str] = field(default_factory=list)


def run_case(case: Case, workdir: Path) -> Outcome:
    from deident import formats, pipeline  # noqa: F401
    from deident.config import default_config
    from deident.entity.pseudonym import Mapping

    root = workdir / case.id.replace(":", "_")
    root.mkdir(parents=True, exist_ok=True)
    source = root / case.filename
    try:
        case.build(source)
    except Exception as exc:   # noqa: BLE001
        return Outcome(case, ERROR, [f"만들기 실패: {type(exc).__name__}: {exc}"[:160]])

    # 심은 값이 정말 들어갔는지부터 확인한다 — 안 들어갔으면 '깨끗함'이 헛것이다
    absent = [v for v in case.must_remove if not residue(source, v)]
    if absent:
        return Outcome(case, ERROR, [f"심기 실패: {len(absent)}개 값이 원본에 없다"])

    config = default_config(root)
    config.workspace.prepare()
    try:
        result = pipeline.apply_file(source, config, Mapping())
    except Exception as exc:   # noqa: BLE001
        return Outcome(case, ERROR, [f"{type(exc).__name__}"])

    if result.blocked:
        return Outcome(case, BLOCKED, [i.code for i in result.issues if i.level == "block"])
    if not result.outputs:
        lines = result.gate.summary_lines() if result.gate else []
        return Outcome(case, BLOCKED, sorted(set(lines)))

    for out in result.outputs:
        broken = opens(out)
        if broken:
            return Outcome(case, ERROR, [f"산출물이 열리지 않는다: {out.suffix} {broken}"])

    left: list[str] = []
    soft_left: list[str] = []
    for out in result.outputs:
        for idx, value in enumerate(case.must_remove):
            for where in residue(out, value):
                left.append(f"{out.suffix} 값{idx + 1} @ {where}")
        for value in case.soft:
            for where in residue(out, value):
                soft_left.append(f"{out.suffix} @ {where}")
    return Outcome(case, LEAK if left else CLEAN, sorted(set(left)), sorted(set(soft_left)))


def available(case: Case) -> bool:
    try:
        if case.fmt == "xlsx":
            import openpyxl  # noqa: F401
        elif case.fmt == "docx":
            import docx  # noqa: F401
        elif case.fmt == "pptx":
            import pptx  # noqa: F401
        elif case.fmt == "pdf":
            import fitz  # noqa: F401
            return korean_font() is not None
        elif case.fmt == "hwpx":
            import lxml  # noqa: F401
    except ImportError:
        return False
    return True


def run_matrix(cases: list[Case] | None = None) -> list[Outcome]:
    workdir = Path(tempfile.mkdtemp(prefix="deident-surfaces-"))
    try:
        return [run_case(c, workdir) for c in (cases or CASES) if available(c)]
    finally:
        shutil.rmtree(workdir, ignore_errors=True)


def main(argv: list[str]) -> int:
    only = set(argv)
    cases = [c for c in CASES if not only or c.fmt in only or c.id in only]
    outcomes = run_matrix(cases)
    counts: dict[str, int] = {}
    for o in outcomes:
        counts[o.status] = counts.get(o.status, 0) + 1
        mark = {CLEAN: "✓", BLOCKED: "■", LEAK: "✗", ERROR: "?"}[o.status]
        soft = f"  (이름 미탐 {len(o.soft_left)}곳)" if o.soft_left else ""
        print(f"{mark} {o.case.fmt:5} {o.case.label or o.case.surface:22} {o.status}{soft}")
        for line in o.detail[:6]:
            print(f"      · {line}")
    print("\n" + " / ".join(f"{k} {v}" for k, v in counts.items()))
    return 0 if set(counts) <= {CLEAN} else 2


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
