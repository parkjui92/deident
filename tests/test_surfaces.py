"""숨은 자리 회귀 시험.

화면에 보이지 않는 자리에 값을 심은 문서를 자리 하나당 하나씩 돌린다(surfaces.py).
통과 기준은 셋이다 — 게이트를 통과하고, 독립 판정기가 값을 찾지 못하고, 산출물이
그 형식의 라이브러리로 다시 열린다.

뒤쪽 시험들은 게이트 자체를 겨눈다. 치환기를 고쳐서 표가 깨끗해져도, 게이트가 그
자리를 못 보면 다음에 생기는 구멍은 조용히 지나간다. 실제로 그랬다 — 시트 이름·
책갈피·PDF 주석에 남은 값을 게이트가 통과시켰다.
"""

from __future__ import annotations

import io
import shutil
import sys
import zipfile
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import surfaces  # noqa: E402
from deident import formats, guard, pipeline  # noqa: E402,F401
from deident.apply.xml_apply import ident_safe  # noqa: E402
from deident.config import default_config  # noqa: E402
from deident.entity.pseudonym import Mapping  # noqa: E402
from deident.extract import hidden  # noqa: E402
from deident.model import CAT_PERSON, KIND_HIDDEN  # noqa: E402
from deident.verify.gate import run_gate  # noqa: E402
from surfaces import HIDDEN_PHONE, KNOWN, NOTE  # noqa: E402


def _case(case_id: str) -> surfaces.Case:
    case = next(c for c in surfaces.CASES if c.id == case_id)
    if not surfaces.available(case):
        pytest.skip(f"{case.fmt} 처리에 필요한 라이브러리(또는 한글 글꼴)가 없다")
    return case


def _apply(case_id: str, tmp_path: Path):
    case = _case(case_id)
    source = tmp_path / case.filename
    case.build(source)
    config = default_config(tmp_path)
    config.workspace.prepare()
    return source, config, pipeline.apply_file(source, config, Mapping())


# ── 자리 하나당 문서 하나 ───────────────────────────────────────────────────
@pytest.mark.parametrize("case", surfaces.CASES, ids=lambda c: c.id)
def test_hidden_surface_is_cleaned(case: surfaces.Case, tmp_path: Path):
    if not surfaces.available(case):
        pytest.skip(f"{case.fmt} 처리에 필요한 라이브러리(또는 한글 글꼴)가 없다")
    outcome = surfaces.run_case(case, tmp_path)
    assert outcome.status == surfaces.CLEAN, f"{case.label}: {outcome.status} {outcome.detail}"


# ── 게이트가 스스로 본다 ────────────────────────────────────────────────────
def _gate_on_untouched(case_id: str, tmp_path: Path, mapping: Mapping):
    """치환을 거치지 않은 파일을 그대로 게이트에 넣는다 — 치환기가 놓쳤을 때의 모습."""
    case = _case(case_id)
    config = default_config(tmp_path)
    config.workspace.prepare()
    leaked = config.workspace.out / case.filename
    case.build(leaked)
    return run_gate(leaked, mapping, config)


def test_gate_sees_attribute_values(tmp_path: Path):
    """시트 이름은 글자 노드가 아니라 속성에 산다."""
    mapping = Mapping()
    mapping.pseudonym_for(KNOWN, CAT_PERSON)
    gate = _gate_on_untouched("xlsx:sheet_name", tmp_path, mapping)
    assert not gate.passed
    assert any(part.startswith("xl/workbook.xml") for _, part in gate.leaks)


def test_gate_sees_inside_embedded_workbook(tmp_path: Path):
    """차트에 딸린 엑셀은 압축 안의 압축이다 — 풀어 보지 않으면 보이지 않는다."""
    mapping = Mapping()
    mapping.pseudonym_for(KNOWN, CAT_PERSON)
    gate = _gate_on_untouched("pptx:chart", tmp_path, mapping)
    assert not gate.passed
    assert any(".xlsx>" in part for _, part in gate.leaks)


@pytest.mark.parametrize("case_id", ["pdf:bookmark", "pdf:annot_note", "pdf:link",
                                     "pdf:attachment", "pdf:form_field"])
def test_gate_sees_pdf_side_content(case_id: str, tmp_path: Path):
    """쪽 글자가 아닌 자리 — 매핑이 비어 있어도 번호 모양만으로 걸려야 한다."""
    gate = _gate_on_untouched(case_id, tmp_path, Mapping())
    assert not gate.passed
    assert KNOWN not in "\n".join(gate.summary_lines())
    assert HIDDEN_PHONE not in "\n".join(gate.summary_lines())


@pytest.mark.parametrize("fmt", ["pptx", "docx", "xlsx"])
def test_gate_passes_blank_office_files(fmt: str, tmp_path: Path):
    """속성값까지 보게 되면서 생길 수 있는 오탐을 막는다.

    빈 발표자료의 `creationId val="1840726560"` 같은 식별자를 전화번호로 읽으면
    어떤 문서도 통과하지 못한다(실측).
    """
    config = default_config(tmp_path)
    config.workspace.prepare()
    blank = config.workspace.out / f"blank.{fmt}"
    if fmt == "pptx":
        pptx = pytest.importorskip("pptx")
        prs = pptx.Presentation()
        prs.slides.add_slide(prs.slide_layouts[0])
        prs.save(blank)
    elif fmt == "docx":
        docx = pytest.importorskip("docx")
        document = docx.Document()
        document.add_paragraph("추진 일정은 2026.03.01부터다.")
        document.save(blank)
    else:
        openpyxl = pytest.importorskip("openpyxl")
        wb = openpyxl.Workbook()
        wb.active.append(["항목", "금액"])
        wb.active.append(["재료비", 1250000])
        wb.save(blank)
    gate = run_gate(blank, Mapping(), config)
    assert gate.passed, gate.summary_lines()


# ── 고친 결과의 모양 ────────────────────────────────────────────────────────
def test_sheet_name_gets_bracketless_pseudonym(tmp_path: Path):
    """대괄호가 든 시트 이름은 엑셀이 받지 않는다 ([…] 는 외부 통합 문서 표기)."""
    openpyxl = pytest.importorskip("openpyxl")
    _, _, result = _apply("xlsx:sheet_name", tmp_path)
    out = next(p for p in result.outputs if p.suffix == ".xlsx")
    titles = openpyxl.load_workbook(out).sheetnames
    assert len(titles) == 2
    assert not any(ch in title for title in titles for ch in "[]")
    assert not any(KNOWN in title for title in titles)


def test_defined_name_keeps_its_reference(tmp_path: Path):
    openpyxl = pytest.importorskip("openpyxl")
    _, _, result = _apply("xlsx:defined_name", tmp_path)
    out = next(p for p in result.outputs if p.suffix == ".xlsx")
    names = openpyxl.load_workbook(out).defined_names
    assert len(names) == 1
    (name, defined), = names.items()
    assert KNOWN not in name and "[" not in name
    assert defined.attr_text == "참여연구원!$D$2"


def test_numeric_cell_becomes_text_cell(tmp_path: Path):
    """숫자 셀에 가명을 넣으면 글자 셀로 바꿔야 파일이 손상으로 읽히지 않는다."""
    openpyxl = pytest.importorskip("openpyxl")
    _, _, result = _apply("xlsx:numeric_rrn", tmp_path)
    out = next(p for p in result.outputs if p.suffix == ".xlsx")
    ws = openpyxl.load_workbook(out).active
    assert ws["E1"].value == "주민등록번호"
    assert isinstance(ws["E2"].value, str) and ws["E2"].value.startswith("[")
    assert ws["A2"].value == 1, "민감하지 않은 숫자 셀은 숫자로 남는다"


def test_ordinary_numbers_are_left_alone(tmp_path: Path):
    openpyxl = pytest.importorskip("openpyxl")
    source = tmp_path / "예산.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["항목", "수량", "단가"])
    ws.append(["시약", 12, 48000])
    ws.append(["장비", 1, 2750000])
    wb.save(source)
    config = default_config(tmp_path)
    config.workspace.prepare()
    result = pipeline.apply_file(source, config, Mapping())
    assert result.gate is not None and result.gate.passed
    out = next(p for p in result.outputs if p.suffix == ".xlsx")
    rows = list(openpyxl.load_workbook(out).active.iter_rows(values_only=True))
    assert rows[1] == ("시약", 12, 48000)
    assert rows[2] == ("장비", 1, 2750000)


def test_embedded_workbook_shares_pseudonyms(tmp_path: Path):
    """본문의 이름과 차트 데이터의 이름이 같은 가명을 받는다."""
    openpyxl = pytest.importorskip("openpyxl")
    _, _, result = _apply("pptx:chart", tmp_path)
    out = next(p for p in result.outputs if p.suffix == ".pptx")
    with zipfile.ZipFile(out) as zf:
        member = next(n for n in zf.namelist() if n.endswith(".xlsx"))
        inner = openpyxl.load_workbook(io.BytesIO(zf.read(member)))
        slide = zf.read("ppt/slides/slide1.xml").decode("utf-8")
    cells = [str(v) for row in inner.active.iter_rows(values_only=True) for v in row if v]
    pseudonym = next(c for c in cells if c.startswith("[") and c.endswith("]"))
    assert pseudonym in slide


def test_hidden_text_stays_out_of_markdown(tmp_path: Path):
    """문서 속성·도구 설명은 읽을 본문이 아니다 — AI 입력용 마크다운에 싣지 않는다."""
    _, _, result = _apply("docx:core_props", tmp_path)
    rendered = next(p for p in result.outputs if p.suffix == ".md").read_text("utf-8")
    assert "과장" not in rendered          # NOTE 의 낱말 — 문서 속성에만 있다
    assert "연구책임자" in rendered         # 본문은 그대로 실린다


def test_report_names_where_it_was_found(tmp_path: Path):
    """숨은 자리에서 잡힌 것은 그 자리 이름으로 보고된다."""
    _, _, result = _apply("xlsx:cell_comment", tmp_path)
    parts = {f.part for f in result.findings}
    assert any("comment" in part for part in parts)
    hidden_segments = [s for s in result.segments if s.kind == KIND_HIDDEN]
    assert hidden_segments


# ── 열기 전 점검 ────────────────────────────────────────────────────────────
def test_guard_sees_values_that_exist_only_in_hidden_places(tmp_path: Path):
    """셀에는 아무것도 없고 메모에만 연락처가 있는 명단 — '없음'이라 하면 거짓 안심이다."""
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl.comments import Comment

    source = tmp_path / "일정.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["회의", "일시"])
    ws.append(["착수 보고", "3월 둘째 주"])
    ws["A2"].comment = Comment(NOTE, "작성자")
    wb.save(source)

    config = default_config(tmp_path)
    config.workspace.prepare()
    result = guard.check_paths([source], config)[0]
    assert result.checked
    assert guard.LEVEL_RANK[result.level] >= guard.LEVEL_RANK[guard.LEVEL_WARN]


# ── 낱개 규칙 ───────────────────────────────────────────────────────────────
@pytest.mark.parametrize("value", [
    "1840726560",                                   # 식별자
    "{00000000-0000-0000-0000-000000000001}",       # GUID
    "2026-01-01T00:00:00.000",                      # 시각
    "9144000",                                      # 좌표
])
def test_machine_values_in_attributes_are_not_content(value: str):
    assert not hidden.has_content(value, long_digits=False)


@pytest.mark.parametrize("value", [NOTE, HIDDEN_PHONE, "mailto:a@b.re.kr", KNOWN])
def test_real_content_in_attributes_is_seen(value: str):
    assert hidden.has_content(value, long_digits=False)


def test_ident_safe_strips_what_names_cannot_hold():
    assert ident_safe("[연구자A]") == "연구자A"
    assert ident_safe("[휴대전화1]") == "휴대전화1"


def test_corpus_plants_what_it_claims(tmp_path: Path):
    """심기 실패를 '깨끗함'으로 읽지 않도록 — 원본에서는 값이 반드시 보여야 한다."""
    for case in surfaces.CASES:
        if not surfaces.available(case):
            continue
        source = tmp_path / case.id.replace(":", "_") / case.filename
        source.parent.mkdir(parents=True)
        case.build(source)
        for value in case.must_remove:
            assert surfaces.residue(source, value), f"{case.id}: 심은 값이 원본에 없다"
    shutil.rmtree(tmp_path, ignore_errors=True)
